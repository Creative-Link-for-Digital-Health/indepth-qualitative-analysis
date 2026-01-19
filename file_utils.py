"""
File I/O utilities for loading configs, parsing uploads, and saving results.
"""

import json
import toml
import docx
from io import BytesIO
from datetime import datetime
from pathlib import Path

from constants import SECRETS_PATH, OUTPUT_DIR


def load_api_params():
    """Load API parameters from TOML file."""
    with open(SECRETS_PATH, 'r') as f:
        secrets = toml.load(f)
    return {
        'API_KEY': secrets['API_KEY'],
        'API_URL': secrets['API_URL'],
        'MODEL': secrets['MODEL'],
    }


def parse_uploaded_file(uploaded_file) -> str:
    """Parse uploaded file and return text content."""
    filename = uploaded_file.name.lower()

    if filename.endswith('.txt'):
        return uploaded_file.read().decode('utf-8')
    elif filename.endswith('.docx'):
        doc = docx.Document(BytesIO(uploaded_file.read()))
        return '\n'.join([p.text for p in doc.paragraphs])
    else:
        raise ValueError(f"Unsupported file type: {filename}")


def save_results(results: dict, filename: str):
    """Save JSON results to outputs folder."""
    OUTPUT_DIR.mkdir(exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    # Save structured JSON
    output_path = OUTPUT_DIR / f"{timestamp}_{filename}_structured.json"
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2, ensure_ascii=False)

    return output_path


def parse_file_from_path(file_path: Path) -> str:
    """Parse file from filesystem path for batch processing.

    Args:
        file_path: Path object pointing to the transcript file

    Returns:
        str: Text content of the file

    Raises:
        ValueError: If file type is not supported (.txt or .docx)
    """
    with open(file_path, 'rb') as f:
        if file_path.suffix.lower() == '.txt':
            return f.read().decode('utf-8')
        elif file_path.suffix.lower() == '.docx':
            doc = docx.Document(f)
            return '\n'.join([p.text for p in doc.paragraphs])
        else:
            raise ValueError(f"Unsupported file type: {file_path.suffix}")
